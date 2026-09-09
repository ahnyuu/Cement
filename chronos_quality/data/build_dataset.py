"""
Build a long-format, timestamped, per-station time series dataset for the Chronos-2 quality net
from the raw xlsx export (운전 & 품질 데이터_2026.xlsx -- see config.py SOURCE_XLSX).

Output columns: item_id, timestamp, <CONTROL_COLS>, <MONITOR_COLS>, blaine, residue
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config as cfg

OUTPUT_PATH = Path(__file__).resolve().parent / "processed" / "quality_timeseries.csv"


def _find_data_start_row(ws, max_scan_rows: int = 5) -> int:
    """
    1-indexed row where actual data begins in this sheet.

    Normally this is cfg.DATA_START_ROW (header text "근무일자" sits at cfg.HEADER_ROW, e.g. the
    2CM/4CM sheets). The 2026-08-13 export's 3CM sheet has no header row at all -- its data
    starts right at row 1 -- unlike 2CM/4CM in the very same workbook, so we can't just trust the
    global constant for every sheet. Detected per-sheet by scanning for the "근무일자" header
    cell; if it's not found in the first few rows, assume there's no header and data starts at
    row 1.
    """
    for row_num, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True), start=1
    ):
        if row and row[0] == "근무일자":
            return row_num + 1
    return 1


def _load_sheet(ws) -> pd.DataFrame:
    data_start_row = _find_data_start_row(ws)
    rows = list(ws.iter_rows(min_row=data_start_row, values_only=True))
    ncols = len(cfg.STD_COLUMNS_BY_POSITION)
    data = [row[:ncols] for row in rows]
    df = pd.DataFrame(data, columns=cfg.STD_COLUMNS_BY_POSITION)
    return df


def load_raw() -> pd.DataFrame:
    wb = openpyxl.load_workbook(cfg.SOURCE_XLSX, read_only=True, data_only=True)
    frames = []
    for sheet_name in cfg.OPERATING_SHEETS:
        df = _load_sheet(wb[sheet_name])
        df["item_id"] = sheet_name
        frames.append(df)
    return pd.concat(frames, ignore_index=True)


def build_timestamp(df: pd.DataFrame) -> pd.Series:
    date = pd.to_datetime(df["Date"], errors="coerce")
    hour = pd.to_numeric(df["WorkTime"], errors="coerce")
    return date + pd.to_timedelta(hour, unit="h")


def clean(
    df: pd.DataFrame,
    *,
    damper_oob: str = "nan",
    prodtype_mode: str = "drop",
) -> pd.DataFrame:
    """Apply the fixed physical/equipment-limit rules from CLEANING_RULES (config.py).

    Two knobs, both defaulting to the canonical pipeline's behavior; the
    02_physical_outlier_handling/ branch scripts flip exactly one at a time (2026-09-02):

    damper_oob -- damper opening-ratio readings outside [0, 100]%:
      "nan"  (default): every out-of-range value -> NaN.
      "clip" (experiment A): a reading overshooting a bound by <= DAMPER_CLIP_MARGIN_PCT is
             calibration drift -> clipped to the bound (keeps the "fully open / fully closed"
             signal); only readings past that margin -> NaN (sensor fault).

    prodtype_mode -- out-of-scope rows (ProdType != 내수), after per-item_id ffill:
      "drop"        (default): remove those rows entirely.
      "mask_target" (experiment B): keep the rows so their operating-condition covariates stay
                    available as model context, blank only the quality targets for them.
    """
    df = df.copy()

    damper_hi = cfg.CLEANING_RULES["damper_max_pct"]
    damper_lo = cfg.CLEANING_RULES["damper_min_pct"]
    for col in cfg.DAMPER_COLS:
        numeric = pd.to_numeric(df[col], errors="coerce")
        if damper_oob == "nan":
            df.loc[(numeric > damper_hi) | (numeric < damper_lo), col] = np.nan
        elif damper_oob == "clip":
            # conditions reference the original `numeric`, so order is irrelevant (disjoint ranges)
            margin = cfg.DAMPER_CLIP_MARGIN_PCT
            df.loc[(numeric > damper_hi) & (numeric <= damper_hi + margin), col] = damper_hi
            df.loc[(numeric < damper_lo) & (numeric >= damper_lo - margin), col] = damper_lo
            df.loc[(numeric > damper_hi + margin) | (numeric < damper_lo - margin), col] = np.nan
        else:
            raise ValueError(f"damper_oob must be 'nan' or 'clip', got {damper_oob!r}")
    for col in cfg.BAG_FILTER_PRESS_COLS:
        df.loc[pd.to_numeric(df[col], errors="coerce") > cfg.CLEANING_RULES["bag_filter_press_max"], col] = np.nan

    mill_coarse = pd.to_numeric(df["mill_feed_c"], errors="coerce")
    df.loc[mill_coarse > cfg.CLEANING_RULES["mill_feed_c_max"], "mill_feed_c"] = np.nan

    blaine = pd.to_numeric(df["blaine"], errors="coerce")
    df.loc[
        (blaine < cfg.CLEANING_RULES["blaine_min"]) | (blaine > cfg.CLEANING_RULES["blaine_max"]),
        "blaine",
    ] = np.nan

    residue = pd.to_numeric(df["residue"], errors="coerce")
    df.loc[residue > cfg.CLEANING_RULES["residue_max"], "residue"] = np.nan

    # Sensor-fault sentinels (physical-limit, not statistical -- see CLEANING_RULES docstring
    # comment in config.py for the gap analysis behind each cutoff; 2026-08-31 data-quality review).
    gas_temp = pd.to_numeric(df["mill_out_gas_temp"], errors="coerce")
    df.loc[gas_temp > cfg.CLEANING_RULES["mill_out_gas_temp_max"], "mill_out_gas_temp"] = np.nan

    mater_temp = pd.to_numeric(df["mill_out_mater_temp"], errors="coerce")
    df.loc[mater_temp > cfg.CLEANING_RULES["mill_out_mater_temp_max"], "mill_out_mater_temp"] = np.nan

    roller_p2 = pd.to_numeric(df["RP_roller_p2"], errors="coerce")
    df.loc[roller_p2 > cfg.CLEANING_RULES["roller_p2_max"], "RP_roller_p2"] = np.nan

    df = df.sort_values(["item_id", "timestamp"])
    blank_prodtype = df["ProdType"].isna() | (df["ProdType"].astype(str).str.strip() == "")
    df.loc[blank_prodtype, "ProdType"] = np.nan
    df["ProdType"] = df.groupby("item_id")["ProdType"].ffill()

    out_of_scope = df["ProdType"] != cfg.CLEANING_RULES["require_prodtype"]  # NaN counts as out-of-scope
    if prodtype_mode == "drop":
        df = df[~out_of_scope]
    elif prodtype_mode == "mask_target":
        df.loc[out_of_scope, cfg.TARGET_COLS] = np.nan
    else:
        raise ValueError(f"prodtype_mode must be 'drop' or 'mask_target', got {prodtype_mode!r}")
    return df


def remove_iqr_outliers(df: pd.DataFrame, cols: list[str], k: float = cfg.IQR_OUTLIER_K):
    """
    Per item_id, per column: flag values outside [Q1 - k*IQR, Q3 + k*IQR] as outliers -> NaN.
    Statistical complement to the fixed-threshold rules in clean() (강원대_성과발표자료.pdf p.8,
    "IQR & Domain 지식 기반 이상치 제거"). Columns with IQR==0 (e.g. near-constant series) are
    left untouched rather than flagging every deviation as an outlier.
    """
    df = df.copy()
    grouped = df.groupby("item_id")
    removed_counts = {}
    for col in cols:
        q1 = grouped[col].transform(lambda s: s.quantile(0.25))
        q3 = grouped[col].transform(lambda s: s.quantile(0.75))
        iqr = q3 - q1
        lo, hi = q1 - k * iqr, q3 + k * iqr
        outlier = ((df[col] < lo) | (df[col] > hi)) & (iqr > 0)
        outlier = outlier.fillna(False)
        removed_counts[col] = int(outlier.sum())
        df.loc[outlier, col] = np.nan
    return df, removed_counts


PREV_QUALITY_COLS = cfg.PREV_QUALITY_COLS


def add_prev_quality_covariates(df: pd.DataFrame) -> pd.DataFrame:
    """Add previous quality only at scheduled 4-hour measurement rows.

    The original hourly blaine/residue columns are never altered.  For each station and target,
    values observed at 00/04/08/12/16/20 are treated as measurement events; *_prev at each event
    receives the preceding event's value.  All intervening *_prev rows remain NaN.
    """
    df = df.sort_values(["item_id", "timestamp"]).reset_index(drop=True)
    scheduled = df["timestamp"].dt.hour.isin(cfg.QUALITY_MEASUREMENT_HOURS)
    for target in cfg.TARGET_COLS:
        prev_col = f"{target}_prev"
        df[prev_col] = np.nan
        measured = scheduled & df[target].notna()
        df.loc[measured, prev_col] = df.loc[measured].groupby("item_id")[target].shift(1)
    return df


def raw_feature_cols() -> list[str]:
    """CONTROL_COLS + the MONITOR_COLS that actually come from the raw xlsx (i.e. excluding the
    ones derived later in the pipeline: *_prev).
    Centralized here (instead of copy-pasted) so every preprocessing branch script
    (notebooks/*/build_*.py) computes the same list -- see prepare_base()/finalize_and_save().
    """
    return cfg.CONTROL_COLS + [
        c for c in cfg.MONITOR_COLS
        if c not in PREV_QUALITY_COLS
    ]


def prepare_base(*, damper_oob: str = "nan", prodtype_mode: str = "drop") -> pd.DataFrame:
    """
    Steps shared by every preprocessing branch, up to (but not including) IQR outlier removal:
    load raw xlsx -> build timestamp -> clean() -> numeric conversion.

    Branch scripts (e.g. notebooks/01_iqr_outlier_removal/build_with_iqr.py vs
    build_without_iqr.py) should call this instead of reimplementing these steps, so a fix here
    (like the 2026-08-12 ProdType bug fix) automatically applies to every branch, not just the
    one someone happens to be editing.

    damper_oob / prodtype_mode pass straight through to clean() -- see its docstring. Defaults
    reproduce the canonical pipeline exactly; 02_physical_outlier_handling/ overrides one at a time.
    """
    df = load_raw()
    df["timestamp"] = build_timestamp(df)
    df = df.dropna(subset=["timestamp"])

    df = clean(df, damper_oob=damper_oob, prodtype_mode=prodtype_mode)

    cols = raw_feature_cols()
    for col in cols + cfg.TARGET_COLS:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    return df


def finalize_and_save(
    df: pd.DataFrame,
    output_path: Path,
    iqr_removed: dict[str, int] | None = None,
    prev_density: str = "sparse",
) -> pd.DataFrame:
    """
    Steps shared by every preprocessing branch, from right after the (optional) IQR step through
    to saving the final CSV: prev-quality covariates -> column selection -> hourly-grid reindex
    -> save + summary print.

    iqr_removed: pass the dict returned by remove_iqr_outliers() if that step ran, so the summary
    print includes an IQR section; leave as None (e.g. build_without_iqr.py, which never calls
    remove_iqr_outliers) to skip that section instead of printing a misleading "0 removed".

    prev_density -- how the blaine_prev/residue_prev covariates are laid out on the hourly grid
    (Experiment 03, 2026-09-02 -- see 03_prev_density_handling/):
      "sparse" (default): value present only at scheduled measurement rows (add_prev_quality_
               covariates' native output), NaN on every hourly row in between.
      "ffill"  (experiment): after the hourly reindex, forward-fill each *_prev per item_id so
               every row carries the most recent lab reading as an explicit anchor. Leak-safe at
               prediction_length<=4: a scheduled row's *_prev already holds the measurement from
               4h earlier, so forward-filling it only ever exposes readings available >=4h before
               any row it lands on (verified in smoketest_prev_density.py).
    """
    cols = raw_feature_cols()

    df = add_prev_quality_covariates(df)

    keep_cols = ["item_id", "timestamp"] + cols + PREV_QUALITY_COLS + cfg.TARGET_COLS
    df = df[keep_cols]

    df = df.sort_values(["item_id", "timestamp"]).drop_duplicates(subset=["item_id", "timestamp"])
    df = df.reset_index(drop=True)

    # Real plant data has gaps (downtime, missing sensor logs) -> irregular timestamps. Chronos-2
    # requires a regular frequency, so reindex each station onto a complete hourly grid and let
    # missing hours become NaN rows (the model handles NaN context/covariates natively).
    reindexed = []
    for item_id, g in df.groupby("item_id", sort=False):
        g = g.set_index("timestamp").drop(columns=["item_id"])
        full_index = pd.date_range(g.index.min(), g.index.max(), freq="h")
        g = g.reindex(full_index)
        g["item_id"] = item_id
        g.index.name = "timestamp"
        g = g.reset_index()
        reindexed.append(g)
    df = pd.concat(reindexed, ignore_index=True)

    if prev_density == "ffill":
        df[PREV_QUALITY_COLS] = df.groupby("item_id", sort=False)[PREV_QUALITY_COLS].ffill()
    elif prev_density != "sparse":
        raise ValueError(f"prev_density must be 'sparse' or 'ffill', got {prev_density!r}")

    feature_cols = cfg.CONTROL_COLS + cfg.MONITOR_COLS
    df = df[["item_id", "timestamp"] + feature_cols + cfg.TARGET_COLS]

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(output_path, index=False)

    print(f"Saved {len(df)} rows -> {output_path}")
    print(df["item_id"].value_counts())
    print()
    print("Timestamp range per item_id:")
    print(df.groupby("item_id")["timestamp"].agg(["min", "max", "count"]))
    print()
    print("Target null rates:")
    print(df[cfg.TARGET_COLS].isna().mean())
    print()
    print(f"prev-quality null rates (prev_density={prev_density!r}):")
    print(df[cfg.PREV_QUALITY_COLS].isna().mean())

    if iqr_removed is not None:
        print()
        print("IQR outlier removal (values set to NaN), top 15 columns by count:")
        iqr_series = pd.Series(iqr_removed).sort_values(ascending=False)
        print(iqr_series.head(15))
        print(f"total values removed by IQR: {iqr_series.sum()}")

    return df


def main() -> None:
    """Default build: full pipeline WITHOUT statistical IQR outlier removal, single canonical
    output path (data/processed/quality_timeseries.csv). Only clean()'s physical/equipment-limit
    thresholds are applied -- statistical (quantile-based) outlier removal was A/B tested
    (2026-08-28) and found to hurt point-forecast accuracy (Chronos-2 masks NaN natively, so
    deleting extreme-but-real values just throws away training signal), so it's no longer the
    default. For side-by-side with/without-IQR comparisons, see notebooks/01_iqr_outlier_removal/
    -- those scripts call prepare_base() / finalize_and_save() directly instead of this main()."""
    df = prepare_base()
    finalize_and_save(df, OUTPUT_PATH, iqr_removed=None)


if __name__ == "__main__":
    main()
